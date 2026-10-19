"""Measured RT sensitivity curves, not a Grade5 carrier material qualification."""
from pathlib import Path
import argparse,csv,hashlib,json,zipfile
import numpy as np
import openpyxl
parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('input',type=Path);parser.add_argument('output',type=Path);args=parser.parse_args();base=args.input;out=args.output;out.mkdir(parents=True,exist_ok=False)
manifest=json.loads((Path(__file__).parent/'sources.json').read_text())
for item in manifest['files']:
 source=base/item['filename'];assert source.stat().st_size==item['bytes'] and hashlib.sha256(source.read_bytes()).hexdigest()==item['sha256']
with zipfile.ZipFile(base/'stress-strain curves.zip') as archive:
 assert set(archive.namelist())==set(manifest['archive_members'])
 for name in manifest['archive_members']:assert (base/'curves'/name).read_bytes()==archive.read(name)
properties=openpyxl.load_workbook(base/'mechanical property table.xlsx',read_only=True,data_only=True);records=[]
# Use the first listed process set, fixed independently of carrier response (275W,800mm/s).
for group,sheet,curve in [('flat set 1','Flat specimen - set 1','f1'),('flat set 2','Flat specimen - set 2','f1'),('round set 1','Round specimen - set 1','r1'),('round set 2','Round specimen - set 2','r1')]:
 data=np.array([[float(a),float(b)] for a,b in openpyxl.load_workbook(base/'curves'/(group+'.xlsx'),read_only=True,data_only=True)[curve].iter_rows(min_row=4,values_only=True) if isinstance(a,(int,float)) and isinstance(b,(int,float))]);prop=next(row for row in properties[sheet].iter_rows(min_row=3,values_only=True) if row[0]==1);uts,yield_stress,neck,fracture,EGPa=prop[1:6];E=EGPa*1000;imax=int(np.argmax(data[:,1]));pre=data[:imax+1];e=pre[:,0];s=pre[:,1];true=s*(1+e);total=np.log1p(e);plastic=total-true/E
 raw=out/(group.replace(' ','-')+'-pre-necking.csv')
 with raw.open('w') as f:
  w=csv.writer(f);w.writerow(['engineering_strain','engineering_stress_MPa','true_strain','true_stress_MPa','true_plastic_strain_using_published_E']);w.writerows(zip(e,s,total,true,plastic))
 # Initial yield for this *diagnostic model* is the measured Rp0.2; reset this model's plastic origin there.
 j=int(np.flatnonzero(s>=yield_stress)[0]);fraction=(yield_stress-s[j-1])/(s[j]-s[j-1]);ey=float(e[j-1]+fraction*(e[j]-e[j-1]));sy=float(yield_stress*(1+ey));ep_y=float(np.log1p(ey)-sy/E);x=plastic[j:]-ep_y;y=true[j:];valid=x>0;x=x[valid];y=y[valid];order=np.argsort(x,kind='stable');x=x[order];y=y[order];assert np.all(np.diff(x)>0)
 # Explicit unweighted isotonic regression (PAVA), no extrapolation or hand-picked curve points.
 x=np.r_[0.,x];y=np.r_[sy,y];blocks=[]
 for i,val in enumerate(y):
  blocks.append([i,i+1,float(val),1])
  while len(blocks)>1 and blocks[-2][2]>blocks[-1][2]:
   right=blocks.pop();left=blocks.pop();weight=left[3]+right[3];blocks.append([left[0],right[1],(left[2]*left[3]+right[2]*right[3])/weight,weight])
 fit=y.copy()
 for a,b,value,weight in blocks:fit[a:b]=value
 assert abs(fit[0]-sy)<1e-8
 law=out/(group.replace(' ','-')+'-plastic.inp')
 with law.open('w') as f:
  f.write('** Diagnostic measured-curve sensitivity: Ti Grade23 ELI, Z, SR670C5h, Luo2023 processset1.\n** Use the published specimen E; Poisson ratio is not supplied by this dataset.\n** Rp0.2 used as model initial yield; plastic origin shifted there.\n** Isotonic PAVA smoothing before engineering UTS only; no post-necking extrapolation.\n*PLASTIC\n')
  for a,b in zip(fit,x):f.write(f'{a:.12g}, {b:.12g}\n')
 records.append({'group':group,'sheet':curve,'parameter_set':1,'published_E_MPa':E,'published_Rp02_MPa':yield_stress,'published_UTS_MPa':uts,'published_elongation_at_necking_percent':neck,'raw_rows':len(data),'pre_necking_rows':len(pre),'pre_necking_endpoint_engineering_strain':float(e[-1]),'model_initial_true_yield_MPa':sy,'model_plastic_origin_raw_true_plastic_strain':ep_y,'plastic_points':len(x),'plastic_strain_max':float(x[-1]),'true_stress_max_MPa':float(fit[-1]),'isotonic_max_adjustment_MPa':float(abs(fit-y).max()),'isotonic_RMS_adjustment_MPa':float(np.sqrt(np.mean((fit-y)**2))),'raw_csv':raw.name,'plastic_input':law.name,'source_xlsx_sha256':hashlib.sha256((base/'curves'/(group+'.xlsx')).read_bytes()).hexdigest(),'raw_csv_sha256':hashlib.sha256(raw.read_bytes()).hexdigest(),'plastic_input_sha256':hashlib.sha256(law.read_bytes()).hexdigest()})
result={'source_record':'https://zenodo.org/records/6587905','article':'https://pmc.ncbi.nlm.nih.gov/articles/PMC9880387/','license':'CC-BY-4.0','attribution':'Qixiang Luo, Lu Yin, TimothyW.Simpson, AllisonM.Beese, Data in Brief46(2023)108911; original raw dataset Zenodo6587905 version10-16-2023','process':{'machine':'3DSystems ProX DMP320','alloy':'LaserForm Ti Gr23(A), Ti6Al4V ELI','laser_power_W':275,'scan_speed_mm_s':800,'layer_um':60,'hatch_um':82,'orientation':'Z, tensile axis vertical build direction','heat_treatment':'670degC5h stress relief,12h Argon furnace cooling','ambient_temperature_K':'room temperature; exact value not supplied','poissons_ratio':None},'model_scope':'Room-temperature sensitivity based on four individual specimens of the first published process set. Not Grade5, not design allowable, no damage, no anisotropic tensor, no post-necking extrapolation. Rp0.2 is a model convention and not a measured proportional limit.','conversion':'Before engineering UTS only: sigma_true=sigma_eng*(1+e_eng); eps_true=ln(1+e_eng); eps_pl=eps_true-sigma_true/E_published. Plastic origin reset to the interpolated measured Rp0.2 point for the model. Sort positive plastic strain, unweighted PAVA stress smoothing, all resulting measured points retained.','records':records,'extractor_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'download_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in base.iterdir() if p.is_file()}};(out/'provenance.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(records,indent=2))
